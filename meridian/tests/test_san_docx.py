"""A3 + A5 (Word): exact occurrence replacement across runs, tables/headers/footers/comments/hidden text,
embedded images with per-occurrence copy-on-write, unsupported payload reporting, independent verification."""
import copy
import hashlib
import io
import zipfile

import docx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml as docx_parse_xml
from lxml import etree

import fixtures_gen as F
from src.sanitize import APPROVED, BLOCKED, REVIEW, Options, Policy, process_document
from src.sanitize.imaging import ocr_available
from src.sanitize.ooxml import collect_image_refs, replace_image_occurrence
from src.sanitize.textunits import W, W_NS

needs_ocr = pytest.mark.skipif(not ocr_available(), reason="tesseract not installed")


def run(src, out, policy=None):
    return process_document(src, out, policy or Policy(ocr=False), Options())


def para_runs(path):
    d = docx.Document(path)
    return [[(r.text, bool(r.bold), bool(r.italic), r.font.size.pt if r.font.size else None) for r in p.runs] for p in d.paragraphs]


def all_xml_text(path):
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith((".xml", ".rels")))


def zip_members(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def check(res, name):
    return next(c for c in res["verification"]["checks"] if c["name"] == name)


# ------------------------------------------------------------------------------ text across runs
def test_text_spanning_three_runs_is_replaced_once_keeping_neighbours_and_formatting(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    res = run(src, out)
    assert res["release"]["status"] == APPROVED
    runs = para_runs(out)
    ssn_para = runs[1]
    assert ssn_para[0] == ("SSN: ", False, False, None)                    # label untouched
    assert ssn_para[1][0] == "[REDACTED:ID]" and ssn_para[1][1] is True     # replacement inherits the FIRST run's formatting (bold)
    assert ssn_para[-1] == (" end of line", False, False, 14.0)             # trailing run keeps its own font size
    assert len(ssn_para) == 3                                               # emptied middle runs are removed, not left as shells
    assert F.SSN not in all_xml_text(out) and "123-45" not in all_xml_text(out)
    assert check(res, "unrelated_content_preserved")["status"] == "passed"


def test_only_the_detected_occurrence_changes_not_every_equal_string(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    run(src, out)
    texts = [p.text for p in docx.Document(out).paragraphs]
    assert f"Invoice #{F.ACCT} is not an account" in texts                  # same digits, different occurrence, untouched
    assert f"Account number: {F.ACCT}" not in texts and "Account number: ******32" in texts


def test_tables_headers_footers_and_properties_are_covered(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    res = run(src, out)
    d = docx.Document(out)
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert "Email: j***@example.com" in cells and "Card: [REDACTED:CARD]" in cells and "Contact" in cells
    assert d.sections[0].header.paragraphs[0].text == "Header Name: [REDACTED:NAME]"
    assert d.sections[0].footer.paragraphs[0].text == "Footer ordinary text"
    assert F.EMAIL not in (d.core_properties.author or "") and d.core_properties.title == "Synthetic statement"
    pay = {p["payload"]: p["outcome"] for p in res["payloads"]}
    assert pay["tables"] == "sanitised" and pay["headers_footers"] == "sanitised" and pay["metadata"] == "sanitised"


def test_unrelated_paragraphs_are_byte_identical_in_xml(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    run(src, out)
    a = docx.Document(src).paragraphs
    b = docx.Document(out).paragraphs
    assert len(a) == len(b)
    for pa, pb in zip(a, b):
        if "SSN" in pa.text or "Account number" in pa.text:
            continue
        assert etree.tostring(pa._p, method="c14n") == etree.tostring(pb._p, method="c14n")


# ------------------------------------------------------------------------------ hidden / alternate / revision content
def _inject(src, mutate):
    d = docx.Document(src)
    mutate(d)
    d.save(src)


def test_hidden_runs_tracked_deletions_comments_textboxes_and_alt_content(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"

    def mutate(d):
        body = d.element.body
        # hidden (vanish) run
        p = d.add_paragraph()
        r = p.add_run(f"hidden SSN {F.SSN}")
        r._r.get_or_add_rPr().append(docx_parse_xml(f'<w:vanish xmlns:w="{W_NS}"/>'))
        # tracked deletion containing an email
        p2 = d.add_paragraph("Kept sentence. ")
        p2._p.append(docx_parse_xml(
            f'<w:del xmlns:w="{W_NS}" w:id="9" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>old {F.EMAIL}</w:delText></w:r></w:del>'))
        # text box through mc:AlternateContent with a fallback copy (both carry the value)
        tb = f'''<w:p xmlns:w="{W_NS}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><w:r>
          <mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><w:txbxContent><w:p><w:r><w:t>Box: {F.EMAIL}</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>
          <mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>Box: {F.EMAIL}</w:t></w:r></w:p></w:txbxContent></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>'''
        body.insert(len(body) - 1, docx_parse_xml(tb))
        # field instruction with a mailto link
        p3 = d.add_paragraph()
        p3._p.append(docx_parse_xml(f'<w:r xmlns:w="{W_NS}"><w:instrText xml:space="preserve"> HYPERLINK "mailto:{F.EMAIL}" </w:instrText></w:r>'))
        # comment (python-docx 1.2 API)
        d.add_comment(runs=d.paragraphs[0].runs[0] if d.paragraphs[0].runs else p.runs[0], text=f"call {F.PHONE}", author="Reviewer Bob")

    F.make_docx(src)
    _inject(src, mutate)
    res = run(src, out)
    assert res["release"]["status"] == APPROVED, [c for c in res["verification"]["checks"] if c["status"] != "passed"]
    xml = all_xml_text(out)
    for secret in (F.SSN, F.EMAIL, F.PHONE):
        assert secret not in xml
    assert "Kept sentence." in xml and "Reviewer Bob" in xml                  # unrelated text and comment author preserved
    pay = {p["payload"]: p for p in res["payloads"]}
    assert pay["hidden_text"]["outcome"] == "sanitised" and pay["comments"]["outcome"] in ("sanitised", "preserved")
    assert any("hidden" in r["reason"] for r in res["regions"])
    assert any(r["source_method"] == "docx_deleted" for r in res["regions"])
    box = [r for r in res["regions"] if r["rule_id"] == "EMAIL.ADDRESS" and r["unit_id"].count(":p") and "document.xml" in r["unit_id"]]
    assert len(box) >= 3                                                      # choice + fallback + deletion each tracked as own occurrence
    assert any(w for w in [res["payloads"]]) and pay["text_boxes"]["count"] >= 1


def test_external_hyperlink_target_is_sanitised(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"

    def mutate(d):
        p = d.add_paragraph()
        rid = d.part.relate_to(f"mailto:{F.EMAIL}", RT.HYPERLINK, is_external=True)
        p._p.append(docx_parse_xml(f'<w:hyperlink xmlns:w="{W_NS}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}"><w:r><w:t>contact us</w:t></w:r></w:hyperlink>'))

    F.make_docx(src)
    _inject(src, mutate)
    res = run(src, out)
    assert res["release"]["status"] == APPROVED
    assert F.EMAIL not in all_xml_text(out) and "contact us" in all_xml_text(out)


def test_package_thumbnail_is_removed_because_it_is_a_preview_of_the_original(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    assert any("thumbnail" in n for n in zip_members(src))
    res = run(src, out)
    assert not any("thumbnail" in n for n in zip_members(out))
    assert next(p for p in res["payloads"] if p["payload"] == "thumbnails")["outcome"] == "sanitised"
    clean = tmp_path / "clean.docx"
    d = docx.Document()
    d.add_paragraph("nothing sensitive here")
    d.save(clean)
    res2 = run(clean, tmp_path / "clean_out.docx")
    assert any("thumbnail" in n for n in zip_members(tmp_path / "clean_out.docx"))        # no findings -> nothing removed
    assert res2["detected"]["regions"] == 0 and res2["release"]["status"] == APPROVED


# ------------------------------------------------------------------------------ images
@needs_ocr
def test_embedded_image_text_is_blanked_and_the_original_bytes_are_gone(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src, with_image=True)
    before = {hashlib.sha256(v).hexdigest() for n, v in zip_members(src).items() if "/media/" in n}
    res = run(src, out, Policy())
    assert res["release"]["status"] == APPROVED, res["verification"]["checks"]
    after = {hashlib.sha256(v).hexdigest() for n, v in zip_members(out).items() if "/media/" in n}
    assert before and not (before & after)                                   # the sensitive original is not hiding in the package
    assert check(res, "embedded_copies_removed")["status"] == "passed" and check(res, "ocr_recheck")["status"] == "passed"
    img_regions = [r for r in res["regions"] if r["space"] == "original_px"]
    assert img_regions and all(r["anchor"].get("occurrence") for r in img_regions)


@needs_ocr
def test_same_image_used_twice_has_each_occurrence_tracked_and_replaced(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src, with_image=True, shared_image=True)
    media_before = [n for n in zip_members(src) if "/media/" in n]
    assert len(media_before) == 1                                            # python-docx stores the shared image once
    res = run(src, out, Policy())
    assert res["release"]["status"] == APPROVED
    occ = {r["anchor"]["occurrence"] for r in res["regions"] if r["space"] == "original_px"}
    assert len(occ) == 2                                                     # two occurrences, two sets of regions
    ids = [r["region_id"] for r in res["regions"] if r["space"] == "original_px"]
    assert len(ids) == len(set(ids))


def test_copy_on_write_changes_only_the_selected_occurrence(tmp_path):
    """replace_image_occurrence(): the other occurrence of the shared image keeps the ORIGINAL part."""
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src, with_image=True, shared_image=True)
    d = docx.Document(src)
    parts = list(d.part.package.iter_parts())
    roots = {str(p.partname): p._element for p in parts if getattr(p, "_element", None) is not None}
    refs, _ = collect_image_refs(parts, roots)
    assert len(refs) == 2 and refs[0].target is refs[1].target
    original = refs[0].target.blob
    marker = F.image_bytes(__import__("PIL.Image", fromlist=["Image"]).new("RGB", (30, 30), (255, 0, 0)))
    replace_image_occurrence(d.part.package, refs[0], marker, RT.IMAGE)
    d.save(out)
    d2 = docx.Document(out)
    parts2 = list(d2.part.package.iter_parts())
    refs2, _ = collect_image_refs(parts2, {str(p.partname): p._element for p in parts2 if getattr(p, "_element", None) is not None})
    blobs = [r.target.blob for r in refs2]
    assert blobs[0] == marker and blobs[1] == original                       # first changed, second untouched
    assert refs2[0].target is not refs2[1].target


def test_unsupported_embedded_objects_are_reported_and_force_review(tmp_path):
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    d = docx.Document(src)
    ole = Part(PackURI("/word/embeddings/oleObject1.bin"), "application/vnd.openxmlformats-officedocument.oleObject",
               b"\xd0\xcf\x11\xe0 pretend OLE payload", d.part.package)
    d.part.relate_to(ole, RT.OLE_OBJECT)
    d.save(src)
    res = run(src, out)
    assert any(p["payload"] == "embedded_objects" and p["outcome"] == "unsupported" for p in res["payloads"])
    assert res["release"]["status"] == REVIEW and res["incomplete"] and res["unsupported"]
    assert not out.exists() and (tmp_path / "out.docx.quarantined").exists()            # held back, not released
    lax = run(src, tmp_path / "lax.docx", Policy(ocr=False, require_complete_inspection=False))
    assert lax["release"]["status"] == APPROVED and lax["unsupported"]                    # explicit policy choice, still listed


def test_hard_block_value_inside_an_unsupported_object_blocks_release(tmp_path):
    """The same SSN string in the body AND inside an object we cannot open must not slip through."""
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    d = docx.Document(src)
    d.part.relate_to(Part(PackURI("/word/embeddings/oleObject1.bin"), "application/vnd.openxmlformats-officedocument.oleObject",
                          f"\xd0\xcf OLE with {F.SSN} inside".encode("latin-1"), d.part.package), RT.OLE_OBJECT)
    d.save(src)
    res = run(src, out)
    assert res["release"]["status"] == BLOCKED and check(res, "hard_block_values_absent")["status"] == "failed"


def test_unreferenced_zip_members_are_reported_when_dropped(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/orphan_secret.txt", f"leftover {F.SSN}")
    (tmp_path / "orph.docx").write_bytes(buf.getvalue())
    res = run(tmp_path / "orph.docx", out)
    assert any("orphan_package_members_dropped" in w for w in res["warnings"])
    assert "orphan_secret" not in " ".join(zip_members(out)) and F.SSN not in all_xml_text(out)


def test_altchunk_and_ole_markup_are_flagged(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"

    def mutate(d):
        d.element.body.insert(0, docx_parse_xml(f'<w:altChunk xmlns:w="{W_NS}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdX"/>'))

    F.make_docx(src)
    _inject(src, mutate)
    res = run(src, out)
    assert any(u["payload"] == "embedded_objects" for u in res["unsupported"]) and res["release"]["status"] == REVIEW


# ------------------------------------------------------------------------------ verification catches a broken edit
def test_verification_blocks_release_if_an_edit_does_not_happen(tmp_path, monkeypatch):
    from src.sanitize.textunits import TextUnit
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    monkeypatch.setattr(TextUnit, "replace_span", lambda self, s, e, r: True)       # claims success, changes nothing
    res = run(src, out)
    assert res["release"]["status"] == BLOCKED and not out.exists()
    assert check(res, "target_text_removed_exact_occurrences")["status"] == "failed"
    assert check(res, "hard_block_values_absent")["status"] == "failed"
    assert res["sanitisation"]["failed"] > 0


def test_formatting_regressions_in_untouched_text_are_detected(tmp_path, monkeypatch):
    """If the sanitiser disturbed run properties of untouched text, the digest comparison must flag it."""
    from src.sanitize import verify as V
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    F.make_docx(src)
    res_ok = run(src, out)
    assert check(res_ok, "unrelated_content_preserved")["status"] == "passed"
    d = docx.Document(out)
    d.paragraphs[0].runs[0].bold = True                                               # tamper with an untouched paragraph
    tampered = tmp_path / "tampered.docx"
    d.save(tampered)
    from src.sanitize import verify_document
    bad = verify_document(src, tampered, Policy(ocr=False))
    assert check(bad, "unrelated_content_preserved")["status"] == "failed"
